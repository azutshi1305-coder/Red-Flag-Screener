import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

# Make the project root importable so `src` can be found when pytest runs from anywhere.
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.export_excel import FINAL_SCORES_PATH, OUTPUT_PATH, main  # noqa: E402

EXPECTED_SHEETS = [
    "Overview",
    "Watchlist",
    "Score details",
    "Flag details",
    "Key ratios",
    "Methodology",
]

# Row 1 = title, row 2 = subtitle, row 3 = blank, row 4 = table header, row 5+ = data.
WATCHLIST_HEADER_ROW = 4


@pytest.fixture(scope="module")
def workbook():
    """Build the real workbook once, from the real processed data, and load it back."""
    main()
    return load_workbook(OUTPUT_PATH)


@pytest.fixture(scope="module")
def watchlist_rows(workbook):
    """Watchlist sheet as a list of dicts, one per data row, keyed by header text."""
    ws = workbook["Watchlist"]
    headers = [cell.value for cell in ws[WATCHLIST_HEADER_ROW]]
    rows = []
    for row in ws.iter_rows(min_row=WATCHLIST_HEADER_ROW + 1, values_only=True):
        rows.append(dict(zip(headers, row)))
    return rows


def test_all_six_sheets_exist(workbook):
    assert workbook.sheetnames == EXPECTED_SHEETS


def test_watchlist_has_twenty_company_rows(watchlist_rows):
    assert len(watchlist_rows) == 20


def test_watchlist_is_sorted_by_risk_score_descending(watchlist_rows):
    scores = [row["Risk score"] for row in watchlist_rows]
    assert scores == sorted(scores, reverse=True)


@pytest.mark.parametrize(
    "short_name,ticker",
    [("Piramal", "PPLPHARMA.NS"), ("Cipla", "CIPLA.NS")],
)
def test_watchlist_values_match_final_scores_csv(watchlist_rows, short_name, ticker):
    final_scores = pd.read_csv(FINAL_SCORES_PATH)
    expected = final_scores.loc[final_scores["ticker"] == ticker].iloc[0]

    matches = [row for row in watchlist_rows if row["Company"] == short_name]
    assert len(matches) == 1, f"expected exactly one Watchlist row for {short_name}"
    actual = matches[0]

    assert actual["Ticker"] == ticker
    assert actual["Risk score"] == expected["risk_score"]
    assert actual["Piotroski F"] == expected["piotroski_f"]
    assert actual["Altman Z''"] == pytest.approx(expected["altman_z"])
    assert actual["Beneish M"] == pytest.approx(expected["beneish_m"])
