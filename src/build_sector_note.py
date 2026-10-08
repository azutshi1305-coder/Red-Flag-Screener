"""
Builds the written sector note (reports/sector_note.md) into two polished,
shareable documents: reports/sector_note.pdf and reports/sector_note.docx.

The Markdown file is the single source of truth for the text. This script
only parses it and lays it out — it never changes any financial data,
calculation, or existing output. The one exception is the case-study
summary table, which this script builds from data/processed/final_scores.csv
and data/companies.csv each time it runs, so its numbers are always read
from the real scores, never retyped into the Markdown. Colours (navy
#1F2A44, grey box #F3F4F6) match src/export_excel.py, so the note looks
like part of the same project as the Excel workbook.

Run from the project root with: python src/build_sector_note.py
"""

import re
import sys
from pathlib import Path

import pandas as pd
from PIL import Image as PILImage

from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as pdfcanvas
from reportlab.platypus import (
    HRFlowable,
    Image,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Table,
    TableStyle,
)

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

MD_PATH = PROJECT_ROOT / "reports" / "sector_note.md"
CHARTS_DIR = PROJECT_ROOT / "outputs" / "charts"
PDF_PATH = PROJECT_ROOT / "reports" / "sector_note.pdf"
DOCX_PATH = PROJECT_ROOT / "reports" / "sector_note.docx"
FINAL_SCORES_PATH = PROJECT_ROOT / "data" / "processed" / "final_scores.csv"
COMPANIES_PATH = PROJECT_ROOT / "data" / "companies.csv"

# Same brand colours as the Excel workbook (src/export_excel.py), so the
# PDF/Word note reads as part of the same project, not a separate tool.
NAVY = "1F2A44"
GREY_TEXT = "6B7280"
BOX_FILL = "F3F4F6"
DARK_TEXT = "111827"

FOOTER_TEXT = "Nifty Pharma Red-Flag Screen | Annapurna Zutshi | Page"

# A4 page geometry, used to size content so it fits inside 18mm margins.
PAGE_WIDTH_MM = 210
PAGE_HEIGHT_MM = 297
MARGIN_MM = 18
CONTENT_WIDTH_MM = PAGE_WIDTH_MM - 2 * MARGIN_MM
PAGE_USABLE_HEIGHT_MM = PAGE_HEIGHT_MM - 2 * MARGIN_MM

# No chart may stand taller than this fraction of the usable page height
# (the tall Altman-vs-Piotroski scatter chart is the one this actually
# affects; every other chart is well under the cap at full width already).
MAX_IMAGE_HEIGHT_FRACTION = 0.45
MAX_IMAGE_HEIGHT_MM = MAX_IMAGE_HEIGHT_FRACTION * PAGE_USABLE_HEIGHT_MM

CONTENT_WIDTH = CONTENT_WIDTH_MM * mm  # reportlab works in points


# ---------------------------------------------------------------------------
# Shared step 1: turn the Markdown into one list of typed "blocks".
# Both the PDF builder and the Word builder read this same list, so the
# parsing rules below exist exactly once.
# ---------------------------------------------------------------------------


def parse_markdown(text):
    """
    Turns the fixed-structure Markdown in reports/sector_note.md into a
    list of simple typed blocks, e.g.
    {"type": "heading2", "text": "Summary", "section": "Summary"}.

    A hand-rolled parser is used here instead of a general Markdown
    library. This file's grammar is small and completely fixed (one title,
    a handful of "##" headings, bold case-study titles, "View:" lines,
    "- " bullet lists, "_..._" disclaimer lines, and "[CHART: file.png]" /
    "[TABLE: name]" markers), so a library would still need all of the
    same custom translation written by hand - it would add a dependency
    without removing any work.
    """
    lines = text.split("\n")
    n = len(lines)

    def skip_to_next_nonblank(index):
        while lines[index].strip() == "":
            index += 1
        return index

    # The first three non-blank lines are always the fixed page-1 header:
    # "# Title", "## Subtitle", then a plain byline line. This pattern only
    # ever appears once, at the very top of the file, so it is peeled off
    # here rather than handled by the general heading/paragraph rules below
    # (otherwise the subtitle line, which starts with "##", would be
    # mistaken for a real section heading like "Summary" or "Methodology").
    i = skip_to_next_nonblank(0)
    title_line = lines[i]
    i += 1

    i = skip_to_next_nonblank(i)
    subtitle_line = lines[i]
    i += 1

    i = skip_to_next_nonblank(i)
    byline_line = lines[i]
    i += 1

    blocks = [
        {"type": "title", "text": title_line[2:].strip(), "section": None},
        {"type": "subtitle", "text": subtitle_line[3:].strip(), "section": None},
        {"type": "byline", "text": byline_line.strip(), "section": None},
    ]

    current_section = None
    while i < n:
        line = lines[i].rstrip()

        if line.strip() == "":
            i += 1
            continue

        if line.startswith("## "):
            current_section = line[3:].strip()
            blocks.append({"type": "heading2", "text": current_section, "section": current_section})
            i += 1

        elif line.startswith("[CHART:") and line.endswith("]"):
            filename = line[len("[CHART:"):-1].strip()
            blocks.append({"type": "image", "filename": filename, "section": current_section})
            i += 1

        elif line.startswith("[TABLE:") and line.endswith("]"):
            table_name = line[len("[TABLE:"):-1].strip()
            blocks.append({"type": "table", "name": table_name, "section": current_section})
            i += 1

        elif line.startswith("- "):
            # Gather every consecutive bullet line into one block, so each
            # builder treats "one list" as one visual unit, not N separate
            # fiddly items.
            items = []
            while i < n and lines[i].rstrip().startswith("- "):
                items.append(lines[i].rstrip()[2:].strip())
                i += 1
            blocks.append({"type": "bullet_list", "items": items, "section": current_section})

        elif line.startswith("View:"):
            blocks.append({"type": "view_line", "text": line.strip(), "section": current_section})
            i += 1

        elif line.startswith("_") and line.endswith("_") and len(line) > 1:
            # A small italic closing disclaimer, e.g. "_Not investment
            # advice._" - markdown's usual underscore-italic syntax.
            blocks.append({"type": "disclaimer", "text": line.strip("_"), "section": current_section})
            i += 1

        elif current_section == "Case studies" and re.match(r"^\*\*\d+\.", line):
            # A case-study line looks like:
            #   **1. Piramal Pharma: genuine financial stress (...).** Piramal is ...
            # Only the bold prefix is the "title"; the rest of the line is
            # an ordinary paragraph that continues straight after it.
            match = re.match(r"^\*\*(\d+\.[^*]*)\*\*(.*)$", line)
            title_text = match.group(1).strip()
            rest_text = match.group(2).strip()
            blocks.append({"type": "case_study_title", "text": title_text, "section": current_section})
            if rest_text:
                blocks.append({"type": "paragraph", "text": rest_text, "section": current_section})
            i += 1

        else:
            blocks.append({"type": "paragraph", "text": line.strip(), "section": current_section})
            i += 1

    return blocks


def split_long_lists_after_headings(blocks):
    """
    If a heading is immediately followed by a multi-item bullet list,
    splits that list into "first item" and "remaining items" as two
    separate blocks.

    Both builders keep a heading together with whatever block comes right
    after it, so it is never left alone at a page bottom. Without this
    split, that meant gluing the heading to an entire 4-5 item list - and
    if that whole bundle didn't fit the room left on a page, it jumped to
    the next page as one piece, leaving the first page mostly empty. After
    the split, a heading only has to stay with its first bullet; the rest
    flow normally and can start right where there's room for them.
    """
    result = []
    i = 0
    n = len(blocks)
    while i < n:
        block = blocks[i]
        result.append(block)
        is_heading_before_long_list = (
            block["type"] == "heading2"
            and i + 1 < n
            and blocks[i + 1]["type"] == "bullet_list"
            and len(blocks[i + 1]["items"]) > 1
        )
        if is_heading_before_long_list:
            next_block = blocks[i + 1]
            result.append({"type": "bullet_list", "items": next_block["items"][:1], "section": next_block["section"]})
            result.append({"type": "bullet_list", "items": next_block["items"][1:], "section": next_block["section"]})
            i += 2
        else:
            i += 1
    return result


# ---------------------------------------------------------------------------
# Shared step 2: the case-study summary table's data.
# Built from the real scores, not retyped, so it can never drift out of
# sync with the rest of the screen.
# ---------------------------------------------------------------------------

CASE_STUDY_TICKERS = ["PPLPHARMA.NS", "ZYDUSLIFE.NS", "TORNTPHARM.NS", "AJANTPHARM.NS", "ABBOTINDIA.NS"]

# Why each company is flagged is an analyst's read of the numbers (financial
# stress vs. an acquisition artifact vs. a likely false positive) - a
# judgement call, not something the screener calculates - so unlike every
# other value in the table, this one is fixed text.
CASE_STUDY_VERDICTS = {
    "PPLPHARMA.NS": "Financial stress",
    "ZYDUSLIFE.NS": "Acquisition (leverage up)",
    "TORNTPHARM.NS": "Acquisition artifact",
    "AJANTPHARM.NS": "Working-capital deterioration",
    "ABBOTINDIA.NS": "Likely false positive",
}

CASE_STUDY_TABLE_HEADER = ["Company", "Risk band", "Risk score", "Altman Z''", "Piotroski F", "Beneish M", "Verdict"]
CASE_STUDY_TABLE_COL_WIDTHS_MM = [24, 18, 16, 18, 16, 18, 64]  # sums to CONTENT_WIDTH_MM (174)


def case_study_table_rows():
    """
    Builds the case-study summary table's rows from
    data/processed/final_scores.csv (scores) and data/companies.csv (short
    names) - every number here is read at build time, not hardcoded. Only
    the Verdict column is fixed text (CASE_STUDY_VERDICTS, above).
    """
    scores = pd.read_csv(FINAL_SCORES_PATH).set_index("ticker")
    companies = pd.read_csv(COMPANIES_PATH).set_index("ticker")

    rows = [CASE_STUDY_TABLE_HEADER]
    for ticker in CASE_STUDY_TICKERS:
        score_row = scores.loc[ticker]
        rows.append([
            companies.loc[ticker, "short_name"],
            score_row["risk_band"],
            str(int(score_row["risk_score"])),
            f"{score_row['altman_z']:.2f}",
            str(int(score_row["piotroski_f"])),
            f"{score_row['beneish_m']:.2f}",
            CASE_STUDY_VERDICTS[ticker],
        ])
    return rows


# ---------------------------------------------------------------------------
# PDF builder (reportlab)
# ---------------------------------------------------------------------------


def _pdf_styles():
    navy = colors.HexColor("#" + NAVY)
    grey = colors.HexColor("#" + GREY_TEXT)
    dark = colors.HexColor("#" + DARK_TEXT)

    return {
        "title": ParagraphStyle(
            "title", fontName="Helvetica-Bold", fontSize=20, leading=24,
            textColor=navy, spaceAfter=4,
        ),
        "subtitle": ParagraphStyle(
            "subtitle", fontName="Helvetica", fontSize=12, leading=15,
            textColor=grey, spaceAfter=4,
        ),
        "byline": ParagraphStyle(
            "byline", fontName="Helvetica", fontSize=9, leading=12,
            textColor=grey, spaceAfter=4,
        ),
        "heading2": ParagraphStyle(
            "heading2", fontName="Helvetica-Bold", fontSize=13, leading=16,
            textColor=navy, spaceBefore=8, spaceAfter=2,
        ),
        "body": ParagraphStyle(
            "body", fontName="Helvetica", fontSize=10, leading=13,
            textColor=dark, spaceAfter=6, alignment=TA_JUSTIFY,
        ),
        "bullet": ParagraphStyle(
            "bullet", fontName="Helvetica", fontSize=10, leading=13,
            textColor=dark, leftIndent=10, spaceAfter=3, alignment=TA_JUSTIFY,
        ),
        "view": ParagraphStyle(
            "view", fontName="Helvetica-Oblique", fontSize=10, leading=13,
            textColor=dark, spaceAfter=8, alignment=TA_JUSTIFY,
        ),
        "case_study_title": ParagraphStyle(
            "case_study_title", fontName="Helvetica-Bold", fontSize=10, leading=13,
            textColor=dark, spaceBefore=5, spaceAfter=2, keepWithNext=True,
        ),
        "disclaimer": ParagraphStyle(
            "disclaimer", fontName="Helvetica-Oblique", fontSize=8, leading=10,
            textColor=grey, spaceBefore=10,
        ),
    }


def _pdf_image(path):
    """
    Sizes a chart from its real pixel aspect ratio (Pillow is already a
    project dependency), capping its height at MAX_IMAGE_HEIGHT_MM so a
    tall chart (like 04_altman_vs_piotroski.png) can never dominate a
    page. A capped chart ends up narrower than the full content width, so
    it is centered rather than left-aligned.
    """
    px_width, px_height = PILImage.open(path).size
    width = CONTENT_WIDTH
    height = width * px_height / px_width
    max_height = MAX_IMAGE_HEIGHT_MM * mm
    if height > max_height:
        height = max_height
        width = height * px_width / px_height
    image = Image(str(path), width=width, height=height)
    image.hAlign = "CENTER"
    return image


def _pdf_case_study_table():
    """
    Renders the case-study summary table to match the Excel workbook's
    look: a navy header row (white bold text) and light grey shading on
    the data rows (src/export_excel.py NAVY / FORMULA_FILL).
    """
    rows = case_study_table_rows()
    col_widths = [w * mm for w in CASE_STUDY_TABLE_COL_WIDTHS_MM]
    table = Table(rows, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#" + BOX_FILL)),
        ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#" + DARK_TEXT)),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 0), (5, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def _render_pdf_block(block, styles):
    """Converts one parsed block into a list of reportlab flowables."""
    kind = block["type"]

    if kind == "paragraph":
        return [Paragraph(block["text"], styles["body"])]

    if kind == "bullet_list":
        return [Paragraph("•  " + item, styles["bullet"]) for item in block["items"]]

    if kind == "view_line":
        return [Paragraph(block["text"], styles["view"])]

    if kind == "case_study_title":
        return [Paragraph(block["text"], styles["case_study_title"])]

    if kind == "disclaimer":
        return [Paragraph(block["text"], styles["disclaimer"])]

    if kind == "image":
        return [_pdf_image(CHARTS_DIR / block["filename"])]

    if kind == "table":
        return [_pdf_case_study_table()]

    raise ValueError(f"Unexpected block type in PDF builder: {kind}")


def _pdf_summary_box(flowables):
    """
    Wraps the Summary section's text in a single-cell table so it reads as
    a grey callout box with a navy left border - the same "callout" look
    the Excel workbook uses (src/export_excel.py FORMULA_FILL / NAVY).
    """
    table = Table([[flowables]], colWidths=[CONTENT_WIDTH])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + BOX_FILL)),
        ("LINEBEFORE", (0, 0), (0, -1), 3, colors.HexColor("#" + NAVY)),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    return table


class _NumberedCanvas(pdfcanvas.Canvas):
    """
    Draws "Page X of Y" on every page. reportlab only knows the true total
    page count once the whole document has been laid out, so this is the
    standard two-pass trick: buffer every finished page instead of emitting
    it (showPage), then once save() is called and the real total is known,
    replay each buffered page, draw its footer, and emit it for real.
    """

    def __init__(self, *args, **kwargs):
        pdfcanvas.Canvas.__init__(self, *args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self._draw_footer(total_pages)
            pdfcanvas.Canvas.showPage(self)
        pdfcanvas.Canvas.save(self)

    def _draw_footer(self, total_pages):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#" + GREY_TEXT))
        text = f"{FOOTER_TEXT} {self._pageNumber} of {total_pages}"
        self.drawCentredString(A4[0] / 2, 10 * mm, text)
        self.restoreState()


def build_pdf(blocks):
    """Lays out `blocks` into reports/sector_note.pdf."""
    styles = _pdf_styles()

    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=A4,
        leftMargin=MARGIN_MM * mm, rightMargin=MARGIN_MM * mm,
        topMargin=MARGIN_MM * mm, bottomMargin=MARGIN_MM * mm,
    )

    story = []
    i = 0
    n = len(blocks)
    while i < n:
        block = blocks[i]

        if block["type"] == "title":
            story.append(Paragraph(block["text"], styles["title"]))
            i += 1

        elif block["type"] == "subtitle":
            story.append(Paragraph(block["text"], styles["subtitle"]))
            i += 1

        elif block["type"] == "byline":
            story.append(Paragraph(block["text"], styles["byline"]))
            story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#" + NAVY), spaceAfter=12))
            i += 1

        elif block["type"] == "heading2":
            heading_flowables = [
                Paragraph(block["text"], styles["heading2"]),
                HRFlowable(width="100%", thickness=0.75, color=colors.HexColor("#" + NAVY), spaceAfter=8),
            ]

            if block["text"] == "Summary":
                # Collect the Summary section's text (paragraph + bullet
                # list) into one grey callout box. The heading stays
                # outside the box like every other heading, and the chart
                # that follows the text is placed after the box, not
                # inside it.
                j = i + 1
                summary_flowables = []
                while j < n and blocks[j]["section"] == "Summary" and blocks[j]["type"] in ("paragraph", "bullet_list"):
                    summary_flowables.extend(_render_pdf_block(blocks[j], styles))
                    j += 1
                box = _pdf_summary_box(summary_flowables)
                group = heading_flowables + [box]
                if j < n and blocks[j]["type"] == "image" and blocks[j]["section"] == "Summary":
                    group.extend(_render_pdf_block(blocks[j], styles))
                    j += 1
                story.append(KeepTogether(group))
                i = j
            else:
                # Keep the heading together with just the very next block
                # (its first bullet, its one paragraph, its table, or its
                # chart - split_long_lists_after_headings already cut a
                # long list down to a single first item), so a heading is
                # never left alone at the bottom of a page without forcing
                # an entire section to jump to the next one. If a chart
                # immediately follows that first block (as in "Sector
                # overview"), keep the chart in the group too, so it never
                # ends up on a page by itself with no heading above it.
                group = list(heading_flowables)
                consumed = 0
                if i + 1 < n:
                    group.extend(_render_pdf_block(blocks[i + 1], styles))
                    consumed = 1
                    if (
                        i + 2 < n
                        and blocks[i + 2]["type"] == "image"
                        and blocks[i + 2]["section"] == block["section"]
                    ):
                        group.extend(_render_pdf_block(blocks[i + 2], styles))
                        consumed = 2
                story.append(KeepTogether(group))
                i += 1 + consumed

        else:
            story.extend(_render_pdf_block(block, styles))
            i += 1

    doc.build(story, canvasmaker=_NumberedCanvas)


# ---------------------------------------------------------------------------
# Word builder (python-docx)
# ---------------------------------------------------------------------------


def _add_bottom_border(paragraph, color_hex):
    """
    Draws a thin rule under a paragraph (used under the byline and every
    heading). python-docx has no high-level API for paragraph borders, so
    this builds the raw Word XML (<w:pBdr><w:bottom>) directly.
    """
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color_hex)
    borders.append(bottom)
    p_pr.append(borders)


def _set_cell_shading(cell, color_hex):
    """Shades a table cell. No high-level API for this in python-docx."""
    cell_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), color_hex)
    cell_pr.append(shading)


def _set_cell_left_border(cell, color_hex):
    """Adds a thick coloured left border to a cell, for the callout-box look."""
    cell_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), color_hex)
    borders.append(left)
    cell_pr.append(borders)


def _add_field(paragraph, instruction, fallback_text, size=None, color_hex=None):
    """
    Inserts a Word field code (e.g. PAGE or NUMPAGES), used for "Page X of
    Y" in the footer. python-docx has no API for field codes, so this
    builds the raw OOXML directly. Word/LibreOffice compute the real value
    when the document is opened or printed; `fallback_text` is only what
    shows before that first recalculation.
    """
    run = paragraph.add_run()
    if size is not None:
        run.font.size = size
    if color_hex is not None:
        run.font.color.rgb = RGBColor.from_string(color_hex)

    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = f" {instruction} "

    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")

    fallback = OxmlElement("w:t")
    fallback.text = fallback_text

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(fallback)
    run._r.append(fld_char_end)


def _setup_docx_footer(section):
    """
    Every page (including page 1) shows only the standard running footer:
    "Nifty Pharma Red-Flag Screen | Annapurna Zutshi | Page X of Y". The
    byline/source line appears once, under the title, as an ordinary
    paragraph - not repeated in the footer.
    """
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(FOOTER_TEXT + " ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GREY_TEXT)
    _add_field(paragraph, "PAGE", "1", size=Pt(8), color_hex=GREY_TEXT)
    run2 = paragraph.add_run(" of ")
    run2.font.size = Pt(8)
    run2.font.color.rgb = RGBColor.from_string(GREY_TEXT)
    _add_field(paragraph, "NUMPAGES", "1", size=Pt(8), color_hex=GREY_TEXT)


def _add_docx_image(container, path):
    """
    The Word equivalent of _pdf_image(): sizes a chart from its real pixel
    aspect ratio, capped at MAX_IMAGE_HEIGHT_MM, and centers it (a capped
    chart is narrower than the full content width).
    """
    px_width, px_height = PILImage.open(path).size
    width_mm = CONTENT_WIDTH_MM
    height_mm = width_mm * px_height / px_width
    if height_mm > MAX_IMAGE_HEIGHT_MM:
        height_mm = MAX_IMAGE_HEIGHT_MM
        width_mm = height_mm * px_width / px_height

    paragraph = container.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Mm(width_mm), height=Mm(height_mm))
    return paragraph


def _add_docx_case_study_table(document, rows):
    """
    The Word equivalent of _pdf_case_study_table(): a navy header row
    (white bold text) and light grey shading on the data rows, matching
    the Excel workbook's styling.
    """
    header, *data_rows = rows
    table = document.add_table(rows=len(rows), cols=len(header))
    table.autofit = False

    for col_index, width_mm in enumerate(CASE_STUDY_TABLE_COL_WIDTHS_MM):
        for row in table.rows:
            row.cells[col_index].width = Mm(width_mm)

    for col_index, text in enumerate(header):
        cell = table.cell(0, col_index)
        _set_cell_shading(cell, NAVY)
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("FFFFFF")

    for row_index, row_values in enumerate(data_rows, start=1):
        for col_index, value in enumerate(row_values):
            cell = table.cell(row_index, col_index)
            _set_cell_shading(cell, BOX_FILL)
            run = cell.paragraphs[0].add_run(str(value))
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(DARK_TEXT)


def _add_docx_block(container, block):
    """
    Adds one parsed block to a python-docx container that has an
    add_paragraph() method - either the Document itself, or a table cell
    (used for the Summary callout box) - so this one function serves both.

    Returns the last paragraph it added (or None for an image/table, which
    aren't paragraphs), so a caller can chain paragraph_format.keep_with_next
    across it when a chart or table needs to stay with the content before it.
    """
    kind = block["type"]

    if kind == "paragraph":
        p = container.add_paragraph(block["text"])
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    elif kind == "bullet_list":
        last_paragraph = None
        for item in block["items"]:
            last_paragraph = container.add_paragraph(item, style="List Bullet")
            last_paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return last_paragraph

    elif kind == "view_line":
        p = container.add_paragraph()
        run = p.add_run(block["text"])
        run.italic = True
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    elif kind == "case_study_title":
        p = container.add_paragraph()
        run = p.add_run(block["text"])
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(DARK_TEXT)
        p.paragraph_format.keep_with_next = True
        return p

    elif kind == "disclaimer":
        p = container.add_paragraph()
        run = p.add_run(block["text"])
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GREY_TEXT)
        return p

    elif kind == "image":
        _add_docx_image(container, CHARTS_DIR / block["filename"])
        return None

    elif kind == "table":
        _add_docx_case_study_table(container, case_study_table_rows())
        return None

    else:
        raise ValueError(f"Unexpected block type in Word builder: {kind}")


def _add_docx_heading(document, text):
    p = document.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    _add_bottom_border(p, NAVY)
    return p


def _add_docx_summary_box(document, summary_blocks):
    """
    Puts the Summary section's text in a single-cell table shaded grey
    with a navy left border - the Word equivalent of the PDF's callout
    box, matching the Excel workbook's callout styling.
    """
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.width = Mm(CONTENT_WIDTH_MM)
    _set_cell_shading(cell, BOX_FILL)
    _set_cell_left_border(cell, NAVY)

    # A new table cell always starts with one empty paragraph; remove it
    # so the box doesn't start with a blank line before the real content.
    empty_paragraph = cell.paragraphs[0]
    empty_paragraph._element.getparent().remove(empty_paragraph._element)

    for block in summary_blocks:
        _add_docx_block(cell, block)


def build_docx(blocks):
    """Lays out `blocks` into reports/sector_note.docx."""
    document = docx.Document()

    section = document.sections[0]
    section.page_width = Mm(PAGE_WIDTH_MM)
    section.page_height = Mm(PAGE_HEIGHT_MM)
    section.left_margin = Mm(MARGIN_MM)
    section.right_margin = Mm(MARGIN_MM)
    section.top_margin = Mm(MARGIN_MM)
    section.bottom_margin = Mm(MARGIN_MM)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"  # Word has no Helvetica; Arial is the usual substitute
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)

    _setup_docx_footer(section)

    i = 0
    n = len(blocks)
    while i < n:
        block = blocks[i]

        if block["type"] == "title":
            p = document.add_paragraph()
            run = p.add_run(block["text"])
            run.bold = True
            run.font.size = Pt(20)
            run.font.color.rgb = RGBColor.from_string(NAVY)
            i += 1

        elif block["type"] == "subtitle":
            p = document.add_paragraph()
            run = p.add_run(block["text"])
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor.from_string(GREY_TEXT)
            i += 1

        elif block["type"] == "byline":
            p = document.add_paragraph()
            run = p.add_run(block["text"])
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(GREY_TEXT)
            _add_bottom_border(p, NAVY)
            i += 1

        elif block["type"] == "heading2":
            if block["text"] == "Summary":
                _add_docx_heading(document, block["text"])

                j = i + 1
                summary_blocks = []
                while j < n and blocks[j]["section"] == "Summary" and blocks[j]["type"] in ("paragraph", "bullet_list"):
                    summary_blocks.append(blocks[j])
                    j += 1
                _add_docx_summary_box(document, summary_blocks)
                if j < n and blocks[j]["type"] == "image" and blocks[j]["section"] == "Summary":
                    _add_docx_image(document, CHARTS_DIR / blocks[j]["filename"])
                    j += 1
                i = j
            else:
                # Mirror the PDF's grouping: keep_with_next ties the
                # heading to whatever comes right after it (its first
                # bullet, its paragraph, its table, or its chart), and if
                # a chart follows that, ties that content to the chart too
                # - without forcing the rest of a long section to move
                # with it.
                heading_paragraph = _add_docx_heading(document, block["text"])
                if i + 1 < n:
                    heading_paragraph.paragraph_format.keep_with_next = True
                    next_paragraph = _add_docx_block(document, blocks[i + 1])
                    consumed = 1
                    if (
                        i + 2 < n
                        and blocks[i + 2]["type"] == "image"
                        and blocks[i + 2]["section"] == block["section"]
                    ):
                        if next_paragraph is not None:
                            next_paragraph.paragraph_format.keep_with_next = True
                        _add_docx_block(document, blocks[i + 2])
                        consumed = 2
                    i += 1 + consumed
                else:
                    i += 1

        else:
            _add_docx_block(document, block)
            i += 1

    document.save(str(DOCX_PATH))


# ---------------------------------------------------------------------------


def main():
    markdown_text = MD_PATH.read_text(encoding="utf-8")
    blocks = parse_markdown(markdown_text)
    blocks = split_long_lists_after_headings(blocks)

    build_pdf(blocks)
    print(f"Wrote {PDF_PATH}")

    build_docx(blocks)
    print(f"Wrote {DOCX_PATH}")


if __name__ == "__main__":
    main()
