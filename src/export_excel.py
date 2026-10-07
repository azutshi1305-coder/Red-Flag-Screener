"""Builds outputs/red_flag_watchlist.xlsx: the ranked Excel watchlist deliverable, styled as a
polished research workbook.

This script only formats and arranges results that src/scores.py and src/flags.py already
computed - it never recomputes a financial figure, score or flag. The one exception is the
handful of numbers behind the custom flags (receivables growth, revenue growth, inventory-day
change and its sector median) that src/flags.py computes internally but does not write to
final_scores.csv; this script reuses src/flags.py's own functions for those, so the numbers on
the "Flag details" sheet always match the flags shown on the Watchlist.

Run from the project root with: python src/export_excel.py
"""

import sys
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.fill import Blip
from openpyxl.drawing.geometry import PresetGeometry2D
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.picture import PictureFrame
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.properties import PageSetupProperties

PROJECT_ROOT = Path(__file__).resolve().parents[1]
# Lets `from src.flags import ...` work when this file is run as `python src/export_excel.py`.
sys.path.insert(0, str(PROJECT_ROOT))

from src.flags import (  # noqa: E402
    FLAG_COLUMNS,
    build_flag_table,
    inventory_days_change,
    receivables_growth,
    revenue_growth,
    sector_median_inventory_days_change,
)
from src.scores import BENEISH_INDICES, PIOTROSKI_TESTS  # noqa: E402

FINAL_SCORES_PATH = PROJECT_ROOT / "data" / "processed" / "final_scores.csv"
SCORES_PATH = PROJECT_ROOT / "data" / "processed" / "scores.csv"
RATIOS_PATH = PROJECT_ROOT / "data" / "processed" / "ratios.csv"
FINANCIALS_PATH = PROJECT_ROOT / "data" / "processed" / "financials_clean.csv"
COMPANIES_PATH = PROJECT_ROOT / "data" / "companies.csv"
NOTES_PATH = PROJECT_ROOT / "data" / "analyst_notes.csv"
FINDINGS_PATH = PROJECT_ROOT / "data" / "key_findings.txt"
CHART_PATH = PROJECT_ROOT / "outputs" / "charts" / "01_risk_ranking.png"
OUTPUT_PATH = PROJECT_ROOT / "outputs" / "red_flag_watchlist.xlsx"
CHART_ALT_TEXT = (
    "Risk-ranking chart: all 20 Nifty Pharma companies ordered by risk score for FY2026, "
    "showing how the High, Watch and Low risk bands separate across the sector."
)

RUN_DATE = date.today()

# Every table leaves column A as a narrow left margin; the table itself starts at column B.
START_COL = 2

# ---------------------------------------------------------------------------
# Theme: fonts, colors, fills
# ---------------------------------------------------------------------------

FONT_NAME = "Calibri"
NAVY = "1F2A44"
GREY_TEXT = "6B7280"
GREY_NEUTRAL = "9CA3AF"
RED = "D03B3B"
AMBER = "E8A100"
GREEN = "0CA30C"
BORDER_GREY = "E1E4E8"
STRIPE_FILL = "F7F8FA"

BODY_FONT = Font(name=FONT_NAME, size=10.5)
HEADER_FONT = Font(name=FONT_NAME, size=10.5, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FONT = Font(name=FONT_NAME, size=16, bold=True, color=NAVY)
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, italic=True, color=GREY_TEXT)
SECTION_FONT = Font(name=FONT_NAME, size=13, bold=True, color=NAVY)
SECTION_BORDER = Border(bottom=Side(style="thin", color=NAVY))
FORMULA_FONT = Font(name="Consolas", size=10, color="111827")
FORMULA_FILL = PatternFill("solid", fgColor="F3F4F6")

# Readable column headers for the machine-friendly Piotroski test and Beneish index names.
PIOTROSKI_LABELS = {
    "p1_roa_positive": "P1 ROA>0",
    "p2_cfo_positive": "P2 CFO>0",
    "p3_roa_improved": "P3 ROA improved",
    "p4_cfo_above_net_income": "P4 CFO>NI",
    "p5_leverage_fell": "P5 Leverage fell",
    "p6_current_ratio_improved": "P6 Current ratio up",
    "p7_no_new_shares": "P7 No new shares",
    "p8_gross_margin_improved": "P8 Gross margin up",
    "p9_asset_turnover_improved": "P9 Asset turnover up",
}
BENEISH_LABELS = {
    "dsri": "DSRI",
    "gmi": "GMI",
    "aqi": "AQI",
    "sgi": "SGI",
    "depi": "DEPI",
    "sgai": "SGAI",
    "lvgi": "LVGI",
    "tata": "TATA",
}

# A custom Excel number format for interest coverage: anything above 100x reads as ">100x"
# rather than a precise but meaningless multiple; the underlying value is never changed.
INTEREST_COVERAGE_FORMAT = '[>100]">100x";0.00"x"'

# The 11 ratios, in the order they appear in ratios.csv, with a short display label and format.
RATIO_COLUMNS = [
    "current_ratio",
    "debt_to_equity",
    "interest_coverage",
    "roa",
    "roe",
    "gross_margin",
    "ebitda_margin",
    "asset_turnover",
    "receivable_days",
    "inventory_days",
    "cash_conversion",
]
RATIO_LABELS = {
    "current_ratio": "Current ratio",
    "debt_to_equity": "D/E",
    "interest_coverage": "Int. cover",
    "roa": "ROA",
    "roe": "ROE",
    "gross_margin": "Gross margin",
    "ebitda_margin": "EBITDA margin",
    "asset_turnover": "Asset turnover",
    "receivable_days": "Receivable days",
    "inventory_days": "Inventory days",
    "cash_conversion": "Cash conversion",
}
RATIO_FORMATS = {
    "current_ratio": "0.00",
    "debt_to_equity": "0.00",
    "interest_coverage": INTEREST_COVERAGE_FORMAT,
    "roa": "0.0%",
    "roe": "0.0%",
    "gross_margin": "0.0%",
    "ebitda_margin": "0.0%",
    "asset_turnover": "0.00",
    "receivable_days": "0",
    "inventory_days": "0",
    "cash_conversion": "0.00",
}
# Direction that counts as improvement for the "chg" column on Key ratios.
HIGHER_IS_BETTER = {
    "current_ratio",
    "interest_coverage",
    "roa",
    "roe",
    "gross_margin",
    "ebitda_margin",
    "asset_turnover",
    "cash_conversion",
}
LOWER_IS_BETTER = {"debt_to_equity", "receivable_days", "inventory_days"}

# Status pills: bold colored text on a pale fill, used for every risk band / score zone column.
PILL_STYLES = {
    "High": ("FBE3E3", "9B1C1C"),
    "Distress": ("FBE3E3", "9B1C1C"),
    "Likely": ("FBE3E3", "9B1C1C"),
    "Weak": ("FBE3E3", "9B1C1C"),
    "Watch": ("FFF4D6", "8A5A00"),
    "Grey": ("FFF4D6", "8A5A00"),
    "Average": ("F3F4F6", "374151"),
    "Low": ("E3F4E3", "1E6B1E"),
    "Safe": ("E3F4E3", "1E6B1E"),
    "Unlikely": ("E3F4E3", "1E6B1E"),
    "Strong": ("E3F4E3", "1E6B1E"),
}
# Text-bar color for the Watchlist's Risk score column, keyed by risk band.
RISK_BAR_COLORS = {"High": "9B1C1C", "Watch": "8A5A00", "Low": "1E6B1E"}

DOT = "●"  # red dot for a triggered flag
SQUARE = "■"  # text-bar block
CHECK = "✓"
CROSS = "✗"
DASH = "–"
ARROW_RIGHT = "→"

# Short labels for each component of the risk score, in the same order src.flags.score_components
# checks them, so "Key reasons" always lists exactly the flags that make up the risk score.
FLAG_SHORT_LABELS = {
    FLAG_COLUMNS[0]: "Receivables",
    FLAG_COLUMNS[1]: "Cash conversion",
    FLAG_COLUMNS[2]: "Inventory",
    FLAG_COLUMNS[3]: "Debt stress",
    FLAG_COLUMNS[4]: "Low interest cover",
}


# ---------------------------------------------------------------------------
# Low-level styling helpers, shared by every sheet.
# ---------------------------------------------------------------------------


def write_rows(ws, rows, start_row, start_col=START_COL):
    """Write a list of row value-lists starting at (start_row, start_col), in the body font.

    Writing cell-by-cell (rather than ws.append) means rows written earlier by hand - the
    title and subtitle - are never disturbed. Returns the last row written.
    """
    for r_offset, row_values in enumerate(rows):
        row = start_row + r_offset
        for c_offset, value in enumerate(row_values):
            cell = ws.cell(row=row, column=start_col + c_offset, value=value)
            cell.font = BODY_FONT
    return start_row + len(rows) - 1


def write_table(ws, columns, rows, header_row, start_col=START_COL):
    """Write a header row plus data rows, then apply each column's number format and alignment.

    `columns` is a list of (header_text, number_format_or_None, kind) triples, where kind is
    "text" (left-aligned, wrapped), "number" (right-aligned), "code" (centered, for short codes,
    pills, dots and checkmarks) or "bar" (left-aligned, not wrapped, for the text-bar columns).
    Returns the row number of the last data row.
    """
    write_rows(ws, [[header for header, _, _ in columns]], header_row, start_col)
    last_row = write_rows(ws, rows, header_row + 1, start_col)
    for i, (_, number_format, kind) in enumerate(columns):
        col_letter = get_column_letter(start_col + i)
        if number_format is not None:
            apply_number_format(ws, col_letter, number_format, header_row + 1, last_row)
        apply_alignment(ws, col_letter, kind, header_row + 1, last_row)
    return last_row


def apply_alignment(ws, col_letter, kind, first_row, last_row):
    """Align a column per the theme: text left+wrapped, number right, code centered, bar left."""
    horizontal = {"text": "left", "number": "right", "code": "center", "bar": "left"}[kind]
    wrap = kind == "text"
    for row in range(first_row, last_row + 1):
        ws[f"{col_letter}{row}"].alignment = Alignment(horizontal=horizontal, vertical="top" if wrap else "center", wrap_text=wrap)


def apply_number_format(ws, col_letter, number_format, first_row, last_row):
    """Apply an Excel number format string to every cell in one column, rows first_row..last_row."""
    for row in range(first_row, last_row + 1):
        ws[f"{col_letter}{row}"].number_format = number_format


def style_header_row(ws, n_cols, header_row, start_col=START_COL):
    """Navy fill, white bold wrapped text, tall centered header row; autofilter over the table."""
    for i in range(n_cols):
        cell = ws.cell(row=header_row, column=start_col + i)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 34
    first_col_letter = get_column_letter(start_col)
    last_col_letter = get_column_letter(start_col + n_cols - 1)
    ws.auto_filter.ref = f"{first_col_letter}{header_row}:{last_col_letter}{ws.max_row}"


def style_data_rows(ws, n_cols, first_row, last_row, start_col=START_COL):
    """Zebra-stripe alternating rows and add a thin light-grey bottom border to every data row."""
    for row in range(first_row, last_row + 1):
        stripe = (row - first_row) % 2 == 1
        for i in range(n_cols):
            cell = ws.cell(row=row, column=start_col + i)
            if stripe:
                cell.fill = PatternFill("solid", fgColor=STRIPE_FILL)
            existing = cell.border
            cell.border = Border(
                left=existing.left,
                right=existing.right,
                top=existing.top,
                bottom=Side(style="thin", color=BORDER_GREY),
            )


def add_group_separators(ws, col_letters, first_row, last_row):
    """Add a light vertical rule before each column in `col_letters`, without losing bottom borders."""
    for col_letter in col_letters:
        for row in range(first_row, last_row + 1):
            cell = ws[f"{col_letter}{row}"]
            existing = cell.border
            cell.border = Border(
                left=Side(style="thin", color="B8BCC4"),
                right=existing.right,
                top=existing.top,
                bottom=existing.bottom,
            )


def force_vertical_center(ws, n_cols, first_row, last_row, start_col=START_COL):
    """Re-center every cell vertically, keeping its existing horizontal/wrap settings."""
    for row in range(first_row, last_row + 1):
        for i in range(n_cols):
            cell = ws.cell(row=row, column=start_col + i)
            align = cell.alignment
            cell.alignment = Alignment(horizontal=align.horizontal, vertical="center", wrap_text=align.wrap_text)


def apply_pill(ws, col_letter, first_row, last_row):
    """Color each cell by the risk-band/zone word it already holds (bold text, pale fill, centered)."""
    for row in range(first_row, last_row + 1):
        cell = ws[f"{col_letter}{row}"]
        style = PILL_STYLES.get(cell.value)
        if style is None:
            continue
        fill_hex, text_hex = style
        cell.fill = PatternFill("solid", fgColor=fill_hex)
        cell.font = Font(name=FONT_NAME, size=10.5, bold=True, color=text_hex)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def style_dot_column(ws, col_letter, first_row, last_row):
    """Bold red for a triggered flag's dot; leaves blank cells (not triggered) untouched."""
    for row in range(first_row, last_row + 1):
        cell = ws[f"{col_letter}{row}"]
        if cell.value == DOT:
            cell.font = Font(name=FONT_NAME, size=11, bold=True, color=RED)
            cell.alignment = Alignment(horizontal="center", vertical="center")


def style_checkmark_column(ws, col_letter, first_row, last_row):
    """Green check for a passed Piotroski test, red cross for a failed one, grey dash if missing."""
    colors = {CHECK: "1E6B1E", CROSS: "9B1C1C", DASH: GREY_NEUTRAL}
    for row in range(first_row, last_row + 1):
        cell = ws[f"{col_letter}{row}"]
        color = colors.get(cell.value)
        if color is not None:
            cell.font = Font(name=FONT_NAME, size=10.5, bold=True, color=color)
            cell.alignment = Alignment(horizontal="center", vertical="center")


def style_risk_bar_column(ws, bar_col_letter, band_col_letter, first_row, last_row):
    """Color each risk-score text bar by the risk band already shown in the same row."""
    for row in range(first_row, last_row + 1):
        band = ws[f"{band_col_letter}{row}"].value
        color = RISK_BAR_COLORS.get(band)
        cell = ws[f"{bar_col_letter}{row}"]
        if color is not None and cell.value:
            cell.font = Font(name=FONT_NAME, size=10.5, bold=True, color=color)


def style_flags_bar_column(ws, col_letter, first_row, last_row):
    """Color the Total-flags text bar navy, matching the theme's neutral accent color."""
    for row in range(first_row, last_row + 1):
        cell = ws[f"{col_letter}{row}"]
        if cell.value:
            cell.font = Font(name=FONT_NAME, size=10.5, bold=True, color=NAVY)


def wrapped_row_height(texts_and_widths, font_size=10.5, chars_per_width=1.4):
    """Estimate a row height (points) tall enough to show every wrapped text at its column width.

    A conservative rule of thumb (characters per line roughly proportional to column width,
    with a generous line-height buffer), since a slightly tall row is harmless but a clipped
    one hides text. Pass a smaller chars_per_width for monospace (Consolas) text, whose
    characters are wider relative to a column-width unit than the default proportional font.
    """
    max_lines = 1
    for text, width in texts_and_widths:
        if not text:
            continue
        chars_per_line = max(10, int(width * chars_per_width))
        lines = -(-len(str(text)) // chars_per_line)
        max_lines = max(max_lines, lines)
    return max(18, max_lines * (font_size * 1.75))


def write_title_block(ws, title, subtitle):
    """Write the sheet's title (row 1) and one-line subtitle (row 2); row 3 is left blank."""
    cell = ws.cell(row=1, column=START_COL, value=title)
    cell.font = TITLE_FONT
    cell = ws.cell(row=2, column=START_COL, value=subtitle)
    cell.font = SUBTITLE_FONT


def set_column_widths(ws, widths, start_col=START_COL):
    """`widths` is a list of column widths in table order; column A (the margin) is set separately."""
    ws.column_dimensions["A"].width = 2.5
    for i, width in enumerate(widths):
        ws.column_dimensions[get_column_letter(start_col + i)].width = width


def set_print_setup(ws, repeat_header_row=None):
    """Landscape, fit to one page wide, and (for tabular sheets) repeat the header row on every page."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if repeat_header_row is not None:
        ws.print_title_rows = f"{repeat_header_row}:{repeat_header_row}"


def embed_image_with_alt_text(ws, image, cell, alt_text):
    """Embed an image with a custom accessibility description (alt text).

    openpyxl's public add_image() always labels a picture's accessibility description as the
    generic word "Picture"; to set a real description we build the same OneCellAnchor and
    PictureFrame structure openpyxl builds internally, but with our own descr, and attach it to
    the image ourselves before handing it to add_image. Falls back to a plain add_image (with
    the generic default description) if anything about this structure changes in a future
    openpyxl version.
    """
    try:
        row, col = coordinate_to_tuple(cell.upper())
        anchor = OneCellAnchor()
        anchor._from.col = col - 1
        anchor._from.row = row - 1
        anchor.ext = XDRPositiveSize2D(cx=pixels_to_EMU(image.width), cy=pixels_to_EMU(image.height))

        pic = PictureFrame()
        pic.nvPicPr.cNvPr.id = 1
        pic.nvPicPr.cNvPr.name = "Sector summary chart"
        pic.nvPicPr.cNvPr.descr = alt_text
        pic.blipFill.blip = Blip()
        pic.blipFill.blip.cstate = "print"
        pic.spPr.prstGeom = PresetGeometry2D(prst="rect")
        pic.spPr.ln = None
        anchor.pic = pic

        image.anchor = anchor
        ws.add_image(image)
    except Exception:
        ws.add_image(image, cell)


# ---------------------------------------------------------------------------
# Data loading
# ---------------------------------------------------------------------------


def load_data():
    """Load the processed CSVs this workbook presents, joining each company's short name and
    analyst notes onto them.

    The analyst's notes are read straight from data/analyst_notes.csv (not from
    final_scores.csv's own "analyst_note" column), so a note added or edited there always shows
    up here without needing to re-run src/flags.py.
    """
    companies = pd.read_csv(COMPANIES_PATH)
    final_scores = pd.read_csv(FINAL_SCORES_PATH)
    scores = pd.read_csv(SCORES_PATH)
    ratios = pd.read_csv(RATIOS_PATH)
    financials = pd.read_csv(FINANCIALS_PATH)
    notes = pd.read_csv(NOTES_PATH)

    short_names = companies[["ticker", "short_name"]]
    final_scores = final_scores.merge(short_names, on="ticker", how="left")
    scores = scores.merge(short_names, on="ticker", how="left")

    notes_for_merge = notes[["ticker", "note", "short_note"]].rename(
        columns={"note": "note_full", "short_note": "note_short"}
    )
    final_scores = final_scores.merge(notes_for_merge, on="ticker", how="left")
    final_scores["note_full"] = final_scores["note_full"].fillna("")
    final_scores["note_short"] = final_scores["note_short"].fillna("")

    return {
        "companies": companies,
        "final_scores": final_scores,
        "scores": scores,
        "ratios": ratios,
        "financials": financials,
    }


def ordered_tickers(final_scores):
    """Ticker order to use on every sheet: final_scores.csv's own order (risk score descending,
    ties broken by lower Altman Z'' first), so row N means the same company on every sheet.
    """
    return final_scores["ticker"].tolist()


def key_reasons(row):
    """Short "Key reasons" text for the Watchlist: just the flag names, separated by a dot.

    Lists exactly the components that make up the risk score (same order and conditions as
    src.flags.score_components), but as short labels instead of full sentences. The full
    plain-English reasons stay on the Watchlist's "Analyst note"/"Score notes" and on the
    Flag details sheet's "Reasons" column.
    """
    labels = []
    if row["altman_zone"] == "Distress":
        labels.append("Altman Distress")
    elif row["altman_zone"] == "Grey":
        labels.append("Altman Grey")
    if row["beneish_zone"] == "Likely":
        labels.append("Beneish Likely")
    elif row["beneish_zone"] == "Watch":
        labels.append("Beneish Watch")
    if row["piotroski_f"] <= 3:
        labels.append("Weak F")
    for col in FLAG_COLUMNS:
        if row[col] == 1:
            labels.append(FLAG_SHORT_LABELS[col])
    return " · ".join(labels) if labels else "No flags"


def build_flag_detail_table(financials, ratios, final_scores):
    """One row per company with the numbers behind flags F1-F5, plus the full Reasons text.

    Receivables growth, revenue growth, inventory-day change and the sector median change are
    computed with src.flags' own functions on the FY2026 rows, the same way src.flags computes
    the Watchlist flags, so these numbers always agree with each other. Total debt, debt/equity
    and interest coverage are read directly for FY2025 and FY2026. Reasons is the same text
    final_scores.csv already carries for each company.
    """
    table = build_flag_table(financials, ratios)
    fy2026 = table[table["fiscal_year"] == 2026].copy()
    fy2026["receivables_growth"] = receivables_growth(fy2026)
    fy2026["revenue_growth"] = revenue_growth(fy2026)
    fy2026["inventory_days_change"] = inventory_days_change(fy2026)
    fy2026["inventory_days_sector_median_change"] = sector_median_inventory_days_change(fy2026)

    detail = fy2026[
        [
            "ticker",
            "receivables_growth",
            "revenue_growth",
            "cfo_below_net_income_prior2",
            "prior_cfo_below_net_income",
            "cfo_below_net_income",
            "inventory_days_change",
            "inventory_days_sector_median_change",
        ]
    ].rename(
        columns={
            "cfo_below_net_income_prior2": "cfo_below_ni_fy2024",
            "prior_cfo_below_net_income": "cfo_below_ni_fy2025",
            "cfo_below_net_income": "cfo_below_ni_fy2026",
        }
    )

    debt = financials.loc[financials["fiscal_year"].isin([2025, 2026]), ["ticker", "fiscal_year", "total_debt"]]
    debt_wide = debt.pivot(index="ticker", columns="fiscal_year", values="total_debt")
    debt_wide.columns = [f"total_debt_fy{int(year)}" for year in debt_wide.columns]
    debt_wide = debt_wide.reset_index()

    ratio_subset = ratios[ratios["fiscal_year"].isin([2025, 2026])]
    de_wide = ratio_subset.pivot(index="ticker", columns="fiscal_year", values="debt_to_equity")
    de_wide.columns = [f"debt_to_equity_fy{int(year)}" for year in de_wide.columns]
    de_wide = de_wide.reset_index()

    ic_wide = ratio_subset.pivot(index="ticker", columns="fiscal_year", values="interest_coverage")
    ic_wide.columns = [f"interest_coverage_fy{int(year)}" for year in ic_wide.columns]
    ic_wide = ic_wide.reset_index()

    detail = detail.merge(debt_wide, on="ticker", how="left")
    detail = detail.merge(de_wide, on="ticker", how="left")
    detail = detail.merge(ic_wide, on="ticker", how="left")
    detail = detail.merge(final_scores[["ticker", "reasons", "note_full"]], on="ticker", how="left")
    return detail


def compute_headline_numbers(final_scores):
    """Return the sector-wide summary numbers shown in the Overview KPI tiles."""
    band_counts = final_scores["risk_band"].value_counts()
    return {
        "companies_screened": len(final_scores),
        "high_count": int(band_counts.get("High", 0)),
        "watch_count": int(band_counts.get("Watch", 0)),
        "low_count": int(band_counts.get("Low", 0)),
        "beneish_likely_count": int((final_scores["beneish_zone"] == "Likely").sum()),
        "median_altman_z": final_scores["altman_z"].median(),
    }


def ratio_change_cell(name, value_2026, value_2025):
    """Return (display_text, font_color_or_None) for one ratio's FY2026 vs FY2025 change.

    Shows an arrow and the change (percentage points for % ratios, days for day-based ratios,
    otherwise the plain difference), rounded to the same precision the FY2026/FY2025 columns
    display. Green when the (rounded) change points the right way for that ratio (higher is
    better, except Debt/Equity, Receivable days and Inventory days, where lower is better), red
    when it points the wrong way, grey "-" when the change rounds to zero at that precision.
    Interest coverage shows "n/m" (not meaningful) whenever either year is above 100x, since the
    exact gap between two very large multiples is not informative.
    """
    if pd.isna(value_2026) or pd.isna(value_2025):
        return None, None
    if name == "interest_coverage" and (value_2026 > 100 or value_2025 > 100):
        return "n/m", GREY_NEUTRAL

    diff = value_2026 - value_2025
    if RATIO_FORMATS[name] == "0.0%":
        rounded = round(diff * 100, 1)
        sign_format, no_sign_format = "{:+.1f}pp", "{:.1f}pp"
    elif name in ("receivable_days", "inventory_days"):
        rounded = round(diff, 0)
        sign_format, no_sign_format = "{:+.0f}d", "{:.0f}d"
    else:
        rounded = round(diff, 2)
        sign_format, no_sign_format = "{:+.2f}", "{:.2f}"

    if rounded == 0:
        return f"{DASH} {no_sign_format.format(0)}", GREY_NEUTRAL
    magnitude = sign_format.format(rounded)
    arrow = "▲" if rounded > 0 else "▼"
    higher_is_better = name in HIGHER_IS_BETTER
    is_good = (rounded > 0) == higher_is_better
    color = "1E6B1E" if is_good else "9B1C1C"
    return f"{arrow} {magnitude}", color


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------


def build_overview_sheet(wb, data):
    """One-page summary: KPI tiles, key findings, a top-5-risks teaser, contents, and the chart."""
    ws = wb.active
    ws.title = "Overview"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY

    ws.cell(row=1, column=START_COL, value="Nifty Pharma Financial Red-Flag Screener").font = Font(
        name=FONT_NAME, size=22, bold=True, color=NAVY
    )
    subtitle = f"FY2026 | 20 companies | Data: Yahoo Finance, verified against Screener.in | Generated {RUN_DATE.isoformat()}"
    ws.cell(row=2, column=START_COL, value=subtitle).font = Font(name=FONT_NAME, size=11, color=GREY_TEXT)

    _write_kpi_tiles(ws, data)
    next_row = _write_key_findings(ws, start_row=8)
    next_row = _write_top_risks(ws, data, start_row=next_row + 2)
    next_row = _write_contents(ws, start_row=next_row + 2)
    _embed_sector_chart(ws, start_row=next_row + 2)

    set_column_widths(ws, [12] * 20)
    ws.sheet_view.zoomScale = 120
    set_print_setup(ws)
    return ws


def _write_kpi_tiles(ws, data):
    """Six equal-width KPI tiles in one row: a colored top strip, a big number, a label - and a
    single narrow spacer column between each pair of tiles.
    """
    headline = compute_headline_numbers(data["final_scores"])
    tiles = [
        ("Companies screened", headline["companies_screened"], NAVY),
        ("High risk", headline["high_count"], RED),
        ("Watch", headline["watch_count"], AMBER),
        ("Low risk", headline["low_count"], GREEN),
        ('Beneish "Likely"', headline["beneish_likely_count"], NAVY),
        ("Median Altman Z''", round(headline["median_altman_z"], 2), NAVY),
    ]
    strip_row, number_row, label_row = 4, 5, 6
    ws.row_dimensions[strip_row].height = 5
    ws.row_dimensions[number_row].height = 30
    ws.row_dimensions[label_row].height = 16

    tile_width, gap = 2, 1
    for i, (label, value, strip_color) in enumerate(tiles):
        c0 = START_COL + i * (tile_width + gap)
        c1 = c0 + tile_width - 1

        ws.merge_cells(start_row=strip_row, start_column=c0, end_row=strip_row, end_column=c1)
        ws.cell(row=strip_row, column=c0).fill = PatternFill("solid", fgColor=strip_color)

        ws.merge_cells(start_row=number_row, start_column=c0, end_row=number_row, end_column=c1)
        number_cell = ws.cell(row=number_row, column=c0, value=value)
        number_cell.font = Font(name=FONT_NAME, size=22, bold=True, color="111827")
        number_cell.alignment = Alignment(horizontal="center")

        ws.merge_cells(start_row=label_row, start_column=c0, end_row=label_row, end_column=c1)
        label_cell = ws.cell(row=label_row, column=c0, value=label)
        label_cell.font = Font(name=FONT_NAME, size=10, color=GREY_TEXT)
        label_cell.alignment = Alignment(horizontal="center")

        for row in range(strip_row, label_row + 1):
            for col in range(c0, c1 + 1):
                cell = ws.cell(row=row, column=col)
                left = Side(style="thin", color="DDE1E6") if col == c0 else None
                right = Side(style="thin", color="DDE1E6") if col == c1 else None
                top = Side(style="thin", color="DDE1E6") if row == strip_row else None
                bottom = Side(style="thin", color="DDE1E6") if row == label_row else None
                cell.border = Border(left=left, right=right, top=top, bottom=bottom)


def _write_key_findings(ws, start_row):
    """Read data/key_findings.txt and show each line as a wrapped bullet point."""
    ws.cell(row=start_row, column=START_COL, value="Key findings").font = SECTION_FONT
    findings = [line.strip() for line in FINDINGS_PATH.read_text().splitlines() if line.strip()]

    row = start_row + 1
    end_col = START_COL + 13
    for finding in findings:
        display_text = finding.replace("->", ARROW_RIGHT)
        ws.merge_cells(start_row=row, start_column=START_COL, end_row=row, end_column=end_col)
        cell = ws.cell(row=row, column=START_COL, value=f"•  {display_text}")
        cell.font = BODY_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = wrapped_row_height([(display_text, 110)])
        row += 1
    return row


def _write_top_risks(ws, data, start_row):
    """A 5-row teaser table: the five highest-risk companies, with short flag names as the reason.

    Built by hand (not via write_table) because "Key reasons" needs to span several of the
    narrow tile columns to be readable, merged wide, with no autofilter on this small teaser.
    """
    ws.cell(row=start_row, column=START_COL, value="Top 5 risks").font = SECTION_FONT

    rank_col = START_COL
    company_col = START_COL + 1
    band_col = START_COL + 2
    score_col = START_COL + 3
    reasons_col0 = START_COL + 4
    reasons_col1 = reasons_col0 + 9  # merged wide enough for a full flag-name list

    header_row = start_row + 1
    data_start = header_row + 1
    headers = [
        (rank_col, rank_col, "Rank"),
        (company_col, company_col, "Company"),
        (band_col, band_col, "Risk band"),
        (score_col, score_col, "Risk score"),
        (reasons_col0, reasons_col1, "Key reasons"),
    ]
    for c0, c1, text in headers:
        if c1 > c0:
            ws.merge_cells(start_row=header_row, start_column=c0, end_row=header_row, end_column=c1)
        cell = ws.cell(row=header_row, column=c0, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 34

    top5 = data["final_scores"].head(5)
    last_row = data_start + len(top5) - 1
    for offset, (_, row) in enumerate(top5.iterrows()):
        r = data_start + offset
        ws.cell(row=r, column=rank_col, value=offset + 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=company_col, value=row["short_name"]).alignment = Alignment(horizontal="left", vertical="center")
        band_cell = ws.cell(row=r, column=band_col, value=row["risk_band"])
        band_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=score_col, value=int(row["risk_score"])).alignment = Alignment(horizontal="right", vertical="center")
        ws.merge_cells(start_row=r, start_column=reasons_col0, end_row=r, end_column=reasons_col1)
        reasons_cell = ws.cell(row=r, column=reasons_col0, value=key_reasons(row))
        reasons_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        for col in range(rank_col, reasons_col1 + 1):
            ws.cell(row=r, column=col).font = BODY_FONT

    style_data_rows(ws, reasons_col1 - rank_col + 1, data_start, last_row, start_col=rank_col)
    apply_pill(ws, get_column_letter(band_col), data_start, last_row)
    return last_row


def _write_contents(ws, start_row):
    """A clickable link to each other sheet, with a one-line description."""
    ws.cell(row=start_row, column=START_COL, value="Contents").font = SECTION_FONT
    guide = [
        ("Watchlist", "Every company ranked by risk score, with its scores, flags and analyst note."),
        ("Score details", "The FY2026 building blocks behind Altman Z'', Piotroski F and Beneish M."),
        ("Flag details", "The raw numbers behind the five custom red flags."),
        ("Key ratios", "All 11 financial ratios for FY2026 and FY2025, side by side."),
        ("Methodology", "Every formula, threshold and data decision used in this workbook."),
    ]
    row = start_row + 1
    for sheet_name, description in guide:
        link_cell = ws.cell(row=row, column=START_COL, value=sheet_name)
        link_cell.hyperlink = f"#'{sheet_name}'!A1"
        link_cell.font = Font(name=FONT_NAME, size=10.5, color="0563C1", underline="single")
        ws.cell(row=row, column=START_COL + 1, value=description).font = BODY_FONT
        row += 1
    return row


def _embed_sector_chart(ws, start_row):
    """Embed the sector-summary chart, scaled to the page width, with accessibility alt text."""
    image = XLImage(str(CHART_PATH))
    target_width = 820
    scale = target_width / image.width
    image.width = target_width
    image.height = int(image.height * scale)
    cell = f"{get_column_letter(START_COL)}{start_row}"
    embed_image_with_alt_text(ws, image, cell, CHART_ALT_TEXT)


def build_watchlist_sheet(wb, data):
    """Main sheet: one row per company, ranked by risk score (ties broken by lower Altman Z'')."""
    final_scores = data["final_scores"]
    ws = wb.create_sheet("Watchlist")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = RED
    write_title_block(ws, "Watchlist", "Every company ranked by risk score; ties broken by the lower Altman Z''.")

    columns = [
        ("Rank", None, "number"),
        ("Company", None, "text"),
        ("Ticker", None, "code"),
        ("Risk band", None, "code"),
        ("Risk score", None, "number"),
        ("Risk bar", None, "bar"),
        ("Altman Z''", "0.00", "number"),
        ("Altman zone", None, "code"),
        ("Piotroski F", "0", "number"),
        ("Piotroski band", None, "code"),
        ("Beneish M", "0.00", "number"),
        ("Beneish zone", None, "code"),
        ("F1", None, "code"),
        ("F2", None, "code"),
        ("F3", None, "code"),
        ("F4", None, "code"),
        ("F5", None, "code"),
        ("Total flags", None, "number"),
        ("Flags bar", None, "bar"),
        ("Key reasons", None, "text"),
        ("Analyst note", None, "text"),
        ("Score notes", None, "text"),
    ]

    rows = []
    for rank, (_, row) in enumerate(final_scores.iterrows(), start=1):
        flag_values = [row[col] for col in FLAG_COLUMNS]
        flag_cells = [DOT if value == 1 else None for value in flag_values]
        total_flags = sum(1 for value in flag_values if value == 1)
        risk_score = int(row["risk_score"])
        rows.append(
            [
                rank,
                row["short_name"],
                row["ticker"],
                row["risk_band"],
                risk_score,
                SQUARE * risk_score if risk_score > 0 else None,
                row["altman_z"],
                row["altman_zone"],
                row["piotroski_f"],
                row["piotroski_band"],
                row["beneish_m"],
                row["beneish_zone"],
                *flag_cells,
                total_flags,
                SQUARE * total_flags if total_flags > 0 else None,
                key_reasons(row),
                row["note_short"],
                row["score_notes"],
            ]
        )

    header_row = 4
    last_row = write_table(ws, columns, rows, header_row)
    # Every row uses the same comfortable height - the short note always fits on one wrapped line.
    for row_number in range(header_row + 1, last_row + 1):
        ws.row_dimensions[row_number].height = 30

    style_header_row(ws, len(columns), header_row)
    style_data_rows(ws, len(columns), header_row + 1, last_row)

    col = {header: get_column_letter(START_COL + i) for i, (header, _, _) in enumerate(columns)}
    apply_pill(ws, col["Risk band"], header_row + 1, last_row)
    apply_pill(ws, col["Altman zone"], header_row + 1, last_row)
    apply_pill(ws, col["Piotroski band"], header_row + 1, last_row)
    apply_pill(ws, col["Beneish zone"], header_row + 1, last_row)
    for flag in ["F1", "F2", "F3", "F4", "F5"]:
        style_dot_column(ws, col[flag], header_row + 1, last_row)
    style_risk_bar_column(ws, col["Risk bar"], col["Risk band"], header_row + 1, last_row)
    style_flags_bar_column(ws, col["Flags bar"], header_row + 1, last_row)
    add_group_separators(ws, [col["Risk band"], col["F1"], col["Key reasons"]], header_row, last_row)
    force_vertical_center(ws, len(columns), header_row + 1, last_row)

    ws.freeze_panes = f"{col['Ticker']}{header_row + 1}"
    set_column_widths(
        ws,
        [8, 12, 17, 10, 10, 9, 10, 11, 12, 13, 10, 11, 7, 7, 7, 7, 7, 11, 9, 45, 55, 28],
    )
    set_print_setup(ws, repeat_header_row=header_row)
    return ws


def build_score_details_sheet(wb, data):
    """FY2026 components behind each company's Altman Z'', Piotroski F and Beneish M."""
    scores_fy2026 = data["scores"][data["scores"]["fiscal_year"] == 2026]
    ordered = ordered_tickers(data["final_scores"])
    scores_fy2026 = scores_fy2026.set_index("ticker").loc[ordered].reset_index()

    ws = wb.create_sheet("Score details")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREY_TEXT
    write_title_block(ws, "Score details", "The FY2026 inputs behind each company's Altman Z'', Piotroski F and Beneish M.")

    columns = [("Company", None, "text"), ("Ticker", None, "code")]
    columns += [(label, "0.00", "number") for label in ["X1", "X2", "X3", "X4", "Altman Z''"]]
    columns += [(PIOTROSKI_LABELS[t], None, "code") for t in PIOTROSKI_TESTS]
    columns += [("Piotroski F", "0", "number")]
    columns += [(BENEISH_LABELS[name], "0.00", "number") for name in BENEISH_INDICES]
    columns += [("Beneish M", "0.00", "number")]

    rows = []
    for _, row in scores_fy2026.iterrows():
        values = [row["short_name"], row["ticker"], row["x1"], row["x2"], row["x3"], row["x4"], row["altman_z"]]
        for t in PIOTROSKI_TESTS:
            value = row[t]
            values.append(DASH if pd.isna(value) else (CHECK if value == 1 else CROSS))
        values += [row["piotroski_f"]]
        values += [row[f"beneish_{name}"] for name in BENEISH_INDICES]
        values += [row["beneish_m"]]
        rows.append(values)

    # Group header row (above the table header) spanning each score's columns.
    group_row, header_row, data_start = 4, 5, 6
    altman_start = START_COL + 2
    altman_end = altman_start + 4
    piotroski_start = altman_end + 1
    piotroski_end = piotroski_start + len(PIOTROSKI_TESTS)
    beneish_start = piotroski_end + 1
    beneish_end = beneish_start + len(BENEISH_INDICES)
    for label, c0, c1 in [("Altman Z''", altman_start, altman_end), ("Piotroski F", piotroski_start, piotroski_end), ("Beneish M", beneish_start, beneish_end)]:
        ws.merge_cells(start_row=group_row, start_column=c0, end_row=group_row, end_column=c1)
        cell = ws.cell(row=group_row, column=c0, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[group_row].height = 18

    last_row = write_table(ws, columns, rows, header_row)
    style_header_row(ws, len(columns), header_row)
    style_data_rows(ws, len(columns), data_start, last_row)

    for t in PIOTROSKI_TESTS:
        i = [c[0] for c in columns].index(PIOTROSKI_LABELS[t])
        style_checkmark_column(ws, get_column_letter(START_COL + i), data_start, last_row)

    ws.freeze_panes = f"{get_column_letter(START_COL + 1)}{data_start}"
    widths = [14, 17] + [9] * 5 + [11] * 9 + [10] + [9] * 8 + [10]
    set_column_widths(ws, widths)
    set_print_setup(ws, repeat_header_row=header_row)
    return ws


def build_flag_details_sheet(wb, data):
    """Per company: the raw numbers behind flags F1-F5, so a reader can see why a flag did or did not fire."""
    detail = build_flag_detail_table(data["financials"], data["ratios"], data["final_scores"])
    ordered = ordered_tickers(data["final_scores"])
    detail = detail.set_index("ticker").loc[ordered].reset_index()
    short_names = data["companies"].set_index("ticker")["short_name"]

    ws = wb.create_sheet("Flag details")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREY_TEXT
    write_title_block(ws, "Flag details", "The raw numbers behind the five custom red flags F1-F5.")

    columns = [
        ("Company", None, "text"),
        ("Ticker", None, "code"),
        ("Receivables gr. %", "0.0%", "number"),
        ("Revenue gr. %", "0.0%", "number"),
        ("CFO<NI FY24", None, "code"),
        ("CFO<NI FY25", None, "code"),
        ("CFO<NI FY26", None, "code"),
        ("Inv. days chg", "0", "number"),
        ("Sector median chg", "0", "number"),
        ("Total debt FY25", "#,##0", "number"),
        ("Total debt FY26", "#,##0", "number"),
        ("D/E FY25", "0.00", "number"),
        ("D/E FY26", "0.00", "number"),
        ("Int. cover FY25", INTEREST_COVERAGE_FORMAT, "number"),
        ("Int. cover FY26", INTEREST_COVERAGE_FORMAT, "number"),
        ("Reasons", None, "text"),
        ("Full analyst note", None, "text"),
    ]

    def cfo_cell(value):
        return DOT if value == 1 else None

    rows = []
    row_heights = []
    for _, row in detail.iterrows():
        rows.append(
            [
                short_names.loc[row["ticker"]],
                row["ticker"],
                row["receivables_growth"],
                row["revenue_growth"],
                cfo_cell(row["cfo_below_ni_fy2024"]),
                cfo_cell(row["cfo_below_ni_fy2025"]),
                cfo_cell(row["cfo_below_ni_fy2026"]),
                row["inventory_days_change"],
                row["inventory_days_sector_median_change"],
                row.get("total_debt_fy2025"),
                row.get("total_debt_fy2026"),
                row.get("debt_to_equity_fy2025"),
                row.get("debt_to_equity_fy2026"),
                row.get("interest_coverage_fy2025"),
                row.get("interest_coverage_fy2026"),
                row["reasons"],
                row["note_full"],
            ]
        )
        row_heights.append(max(30, wrapped_row_height([(row["reasons"], 70), (row["note_full"], 55)])))

    header_row, data_start = 4, 5
    last_row = write_table(ws, columns, rows, header_row)
    for offset, height in enumerate(row_heights):
        ws.row_dimensions[data_start + offset].height = height

    style_header_row(ws, len(columns), header_row)
    style_data_rows(ws, len(columns), data_start, last_row)
    for cfo_col in ["CFO<NI FY24", "CFO<NI FY25", "CFO<NI FY26"]:
        i = [c[0] for c in columns].index(cfo_col)
        style_dot_column(ws, get_column_letter(START_COL + i), data_start, last_row)
    force_vertical_center(ws, len(columns), data_start, last_row)

    ws.freeze_panes = f"{get_column_letter(START_COL + 1)}{data_start}"
    set_column_widths(ws, [14, 17, 13, 13, 10, 10, 10, 11, 14, 11, 11, 10, 10, 12, 12, 70, 55])
    set_print_setup(ws, repeat_header_row=header_row)
    return ws


def build_key_ratios_sheet(wb, data):
    """The 11 ratios for FY2026 and FY2025, side by side per ratio, with a change column after each."""
    ratios = data["ratios"]
    ordered = ordered_tickers(data["final_scores"])
    short_names = data["companies"].set_index("ticker")["short_name"]
    fy2026 = ratios[ratios["fiscal_year"] == 2026].set_index("ticker")
    fy2025 = ratios[ratios["fiscal_year"] == 2025].set_index("ticker")

    ws = wb.create_sheet("Key ratios")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREY_TEXT
    write_title_block(ws, "Key ratios", "All 11 ratios for FY2026 and FY2025, with the year-on-year change.")

    columns = [("Company", None, "text"), ("Ticker", None, "code")]
    for name in RATIO_COLUMNS:
        fmt = RATIO_FORMATS[name]
        columns.append((f"{RATIO_LABELS[name]} FY26", fmt, "number"))
        columns.append((f"{RATIO_LABELS[name]} FY25", fmt, "number"))
        columns.append((f"{RATIO_LABELS[name]} chg", None, "code"))

    rows = []
    change_positions = []  # (row_offset, col_offset, color) for cells needing a custom font color
    for row_offset, ticker in enumerate(ordered):
        values = [short_names.loc[ticker], ticker]
        col_offset = 2
        for name in RATIO_COLUMNS:
            value_2026 = fy2026.loc[ticker, name] if ticker in fy2026.index else None
            value_2025 = fy2025.loc[ticker, name] if ticker in fy2025.index else None
            values.append(value_2026)
            values.append(value_2025)
            text, color = ratio_change_cell(name, value_2026, value_2025)
            values.append(text)
            if color is not None:
                change_positions.append((row_offset, col_offset + 2, color))
            col_offset += 3
        rows.append(values)

    header_row, data_start = 4, 5
    last_row = write_table(ws, columns, rows, header_row)
    style_header_row(ws, len(columns), header_row)
    style_data_rows(ws, len(columns), data_start, last_row)

    for row_offset, col_offset, color in change_positions:
        cell = ws.cell(row=data_start + row_offset, column=START_COL + col_offset)
        cell.font = Font(name=FONT_NAME, size=10.5, bold=True, color=color)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = f"{get_column_letter(START_COL + 1)}{data_start}"
    set_column_widths(ws, [14, 17] + [11, 11, 10] * len(RATIO_COLUMNS))
    set_print_setup(ws, repeat_header_row=header_row)
    return ws


# Methodology body text: (section_title, [(line_text, is_formula), ...]).
METHODOLOGY_SECTIONS = [
    (
        "Altman Z'' (financial distress)",
        [
            ("X1 = (current assets - current liabilities) / total assets   (liquidity buffer)", True),
            ("X2 = retained earnings / total assets   (cumulated profitability)", True),
            ("X3 = EBIT / total assets   (operating profitability)", True),
            ("X4 = total equity / total liabilities   (equity cushion)", True),
            ("Z'' = 6.56*X1 + 3.26*X2 + 6.72*X3 + 1.05*X4", True),
            ("Zones: Safe if Z'' > 2.6; Grey if 1.1 <= Z'' <= 2.6; Distress if Z'' < 1.1.", True),
        ],
    ),
    (
        "Piotroski F-score (fundamental strength, 0-9)",
        [
            ("1. ROA_t > 0, where ROA_t = net income_t / total assets_(t-1)", True),
            ("2. Operating cash flow_t > 0", True),
            ("3. ROA_t > ROA_(t-1)", True),
            ("4. Operating cash flow_t > net income_t (profits backed by cash)", True),
            (
                "5. Leverage fell: long-term debt / average assets fell year on year. A debt-free "
                "company (zero long-term debt in both years) passes automatically.",
                True,
            ),
            ("6. Current ratio_t > current ratio_(t-1)", True),
            (
                "7. Shares outstanding_t <= shares_(t-1) * 1.01 (increases under 1% are treated as "
                "ESOPs, not an equity raise)",
                True,
            ),
            ("8. Gross margin_t > gross margin_(t-1)", True),
            ("9. Asset turnover_t > asset turnover_(t-1), where turnover_t = revenue_t / assets_(t-1)", True),
            (
                "F = count of tests passed, out of the tests with enough data to run (a test with "
                "missing inputs is left out, never scored as failed).",
                False,
            ),
            ("Bands: Strong (8-9), Average (4-7), Weak (0-3).", True),
        ],
    ),
    (
        "Beneish M-score (earnings-manipulation risk)",
        [
            ("DSRI = (receivables/revenue)_t / (receivables/revenue)_(t-1)", True),
            ("GMI  = gross margin_(t-1) / gross margin_t", True),
            ("AQI  = asset quality_t / asset quality_(t-1), asset quality = 1 - (current assets + net PPE) / total assets", True),
            ("SGI  = revenue_t / revenue_(t-1)", True),
            ("DEPI = depreciation rate_(t-1) / depreciation rate_t, rate = depreciation / (depreciation + net PPE)", True),
            ("SGAI = (SG&A/revenue)_t / (SG&A/revenue)_(t-1)", True),
            ("LVGI = liabilities-to-assets_t / liabilities-to-assets_(t-1), liabilities-to-assets = (current liabilities + long-term debt) / total assets", True),
            ("TATA = (net income_t - operating cash flow_t) / total assets_t", True),
            ("M = -4.84 + 0.920*DSRI + 0.528*GMI + 0.404*AQI + 0.892*SGI + 0.115*DEPI - 0.172*SGAI + 4.679*TATA - 0.327*LVGI", True),
            ("Zones: Likely if M > -1.78; Watch if -2.22 < M <= -1.78; Unlikely if M <= -2.22.", True),
            (
                "The Watch band was added because Ajanta (M=-1.80) and Abbott (M=-1.66) sit close to "
                "the -1.78 cutoff, so a single cutoff would hinge on rounding.",
                False,
            ),
            (
                "Missing inputs default to a neutral value rather than being invented: the seven ratio "
                "indices default to 1.0 (no change), TATA defaults to 0.0 (no accrual). Any company "
                "with a filled index is listed in its Score notes column.",
                False,
            ),
        ],
    ),
    (
        "Custom red flags (F1-F5)",
        [
            (
                "F1 Receivables outpacing sales: flagged when receivables growth exceeds revenue "
                "growth by more than 10 percentage points (FY2026 vs FY2025).",
                False,
            ),
            (
                "F2 Weak cash conversion: flagged when operating cash flow was below net income in at "
                "least 2 of FY2024, FY2025, FY2026 (compares raw rupee values, not a ratio).",
                False,
            ),
            (
                "F3 Inventory buildup: flagged when a company's change in inventory days (FY2026 vs "
                "FY2025) is more than 15 days above the sector median change for the same two years. "
                "This is a relative rule, not a fixed +15 days: the sector median change itself was "
                "+13.46 days, showing a sector-wide buildup, so a fixed threshold would punish "
                "companies merely for following that trend.",
                False,
            ),
            (
                "F4 Debt stress: flagged only when total debt rose AND interest coverage fell AND "
                "total debt exceeds 10% of total equity in FY2026 (a materiality floor so small "
                "borrowings cannot trigger it). Not flagged if interest coverage is missing in either "
                "year.",
                False,
            ),
            (
                "F5 Low interest cover: flagged when interest coverage is below 1.5x. A missing "
                "coverage (no interest expense) is not flagged.",
                False,
            ),
        ],
    ),
    (
        "Risk score and bands",
        [
            (
                "Points: Altman Distress +2, Altman Grey +1; Beneish Likely +2, Watch +1; Piotroski F "
                "<= 3 +1; each triggered custom flag (F1-F5) +1.",
                True,
            ),
            ("Risk score = sum of points. Bands: Low (0-1), Watch (2-3), High (4 or more).", True),
            (
                "Ties on the Watchlist are broken by lower Altman Z'' first (the more distressed "
                "company is ranked higher).",
                False,
            ),
        ],
    ),
    (
        "Data sources",
        [
            (
                "Financial statements: Yahoo Finance (via the yfinance library), FY2026 as the main "
                "year and FY2025 as the secondary year.",
                False,
            ),
            ("Spot-checked against Screener.in for companies with data gaps (see Key decisions below).", False),
            ("All money figures are in Rs crore.", False),
        ],
    ),
    (
        "Key decisions and limitations",
        [
            (
                "Abbott India, Divi's Laboratories and Ajanta Pharma: Yahoo reports no separate "
                "long-term debt, so Total Debt is used in its place for all years, after confirming "
                "the figures match Screener.in's borrowings line.",
                False,
            ),
            (
                "Aurobindo Pharma: FY2025 SG&A is missing from Yahoo. Beneish SGAI is set to neutral "
                "(1.0) for FY2025 and FY2026, and the company is flagged in its score_notes rather "
                "than inventing a number.",
                False,
            ),
            (
                "Dr Reddy's Laboratories: FY2024 depreciation was entered manually from Screener.in's "
                "consolidated P&L (Rs 1,470 crore), because Yahoo did not report it.",
                False,
            ),
            (
                "Torrent Pharmaceuticals: FY2025 and FY2026 retained earnings are rolled forward from "
                "FY2024 using net income and dividends paid (Yahoo figures only), because Screener's "
                "'Reserves' line did not match Yahoo's retained earnings and was rejected as an "
                "inconsistent data source.",
                False,
            ),
            (
                "Limitation: a screener cannot tell acquisition-driven growth from genuine distress. "
                "Torrent's FY2026 borrowings rose sharply (Rs 3,202 crore to Rs 15,026 crore), most "
                "likely funding an acquisition, so its FY2026 scores should be read with that context, "
                "not as organic deterioration. Analyst judgment, not just the scores, is required.",
                False,
            ),
        ],
    ),
]


def build_methodology_sheet(wb):
    """Formulas, thresholds and the key data decisions behind every number in this workbook."""
    ws = wb.create_sheet("Methodology")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    write_title_block(ws, "Methodology", "Every formula, threshold and data decision used in this workbook.")

    row = 4
    for title, lines in METHODOLOGY_SECTIONS:
        heading_cell = ws.cell(row=row, column=START_COL, value=title)
        heading_cell.font = SECTION_FONT
        heading_cell.border = SECTION_BORDER
        row += 1
        for text, is_formula in lines:
            cell = ws.cell(row=row, column=START_COL, value=text)
            if is_formula:
                cell.font = FORMULA_FONT
                cell.fill = FORMULA_FILL
                chars_per_width = 0.95  # Consolas is wider per character than the column-width unit
            else:
                cell.font = BODY_FONT
                chars_per_width = 1.3
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row].height = wrapped_row_height([(text, 120)], chars_per_width=chars_per_width)
            row += 1
        row += 1

    ws.column_dimensions["A"].width = 2.5
    ws.column_dimensions["B"].width = 120
    set_print_setup(ws)
    return ws


def main():
    data = load_data()
    wb = Workbook()
    build_overview_sheet(wb, data)
    build_watchlist_sheet(wb, data)
    build_score_details_sheet(wb, data)
    build_flag_details_sheet(wb, data)
    build_key_ratios_sheet(wb, data)
    build_methodology_sheet(wb)
    wb.active = 0

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"Saved workbook to {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
